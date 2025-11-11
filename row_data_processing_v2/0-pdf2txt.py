#!/usr/bin/env python3
import argparse
import os
import subprocess
import tempfile
from pathlib import Path
from typing import Iterable

import pdfplumber
import docx


def _maybe_two_columns_split_x(page):
    """
    Heuristic: look at character x-positions and find the largest horizontal gap.
    If it's big enough (>= 15% of page width), treat it as a column gutter and return split x.
    """
    xs = sorted({c["x0"] for c in page.chars})  # unique-ish x positions
    if len(xs) < 10:
        return None
    gaps = [(xs[i + 1] - xs[i], i) for i in range(len(xs) - 1)]
    max_gap, idx = max(gaps, key=lambda t: t[0])
    if max_gap >= page.width * 0.15:
        return xs[idx] + max_gap / 2.0
    return None


def extract_pdf_text(pdf_path: str) -> str:
    """Extract text from a PDF using pdfplumber with basic two-column handling."""
    text_chunks: list[str] = []
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            # Try to detect a two-column layout regardless of orientation
            split_x = _maybe_two_columns_split_x(page)
            if split_x:
                left = page.within_bbox((0, 0, split_x, page.height)).extract_text(
                    x_tolerance=3, y_tolerance=3, keep_blank_chars=False, use_text_flow=True
                ) or ""
                right = page.within_bbox((split_x, 0, page.width, page.height)).extract_text(
                    x_tolerance=3, y_tolerance=3, keep_blank_chars=False, use_text_flow=True
                ) or ""
                text_chunks.extend([left, right])
            else:
                txt = page.extract_text(
                    x_tolerance=3, y_tolerance=3, keep_blank_chars=False, use_text_flow=True
                ) or ""
                text_chunks.append(txt)
    return "\n\n".join(chunk for chunk in text_chunks if chunk.strip())


def extract_docx_text(docx_path: str) -> str:
    """Extract text from a DOCX file, including paragraphs, tables, headers, and footers."""
    d = docx.Document(docx_path)
    parts: list[str] = []

    # Body paragraphs
    parts.extend(p.text for p in d.paragraphs)

    # Tables
    for table in d.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))

    # Headers & footers
    for section in d.sections:
        if section.header:
            parts.extend(p.text for p in section.header.paragraphs)
        if section.footer:
            parts.extend(p.text for p in section.footer.paragraphs)

    return "\n".join(s for s in parts if s and s.strip())


def extract_doc_text(doc_path: str) -> str:
    """
    Extract text from a legacy DOC using LibreOffice headless conversion.
    Forces UTF-8 via the 'Text (encoded)' filter when available.
    """
    with tempfile.TemporaryDirectory() as tmpdir:
        # Some builds use this exact syntax; others accept without the ":UTF8,LF" tail.
        # If it fails, fall back to plain 'txt:Text'.
        cmd_utf8 = [
            "soffice", "--headless", "--convert-to",
            'txt:Text (encoded):UTF8,LF', "--outdir", tmpdir, doc_path
        ]
        cmd_plain = [
            "soffice", "--headless", "--convert-to",
            'txt:Text', "--outdir", tmpdir, doc_path
        ]
        try:
            subprocess.run(cmd_utf8, check=True, capture_output=True)
        except subprocess.CalledProcessError:
            subprocess.run(cmd_plain, check=True, capture_output=True)

        txt_name = Path(doc_path).with_suffix(".txt").name
        txt_path = Path(tmpdir) / txt_name
        return txt_path.read_text(encoding="utf-8", errors="ignore")


def iter_files(root: Path) -> Iterable[Path]:
    for p in root.rglob("*"):
        if p.is_file():
            yield p


def convert_files_to_text(input_dir: str, output_dir: str) -> None:
    in_root = Path(input_dir)
    out_root = Path(output_dir)
    out_root.mkdir(parents=True, exist_ok=True)

    for src in iter_files(in_root):
        ext = src.suffix.lower()
        if ext not in {".pdf", ".docx", ".doc"}:
            continue

        rel = src.relative_to(in_root)
        dest = (out_root / rel).with_suffix(".txt")
        dest.parent.mkdir(parents=True, exist_ok=True)

        try:
            if ext == ".pdf":
                text = extract_pdf_text(str(src))
                # Optional: if text is nearly empty, the PDF might be scanned.
                # Consider OCR pre-pass: `ocrmypdf --sidecar sidecar.txt src.pdf ocr.pdf`
            elif ext == ".docx":
                text = extract_docx_text(str(src))
            else:  # .doc
                text = extract_doc_text(str(src))

            dest.write_text(text or "", encoding="utf-8")
            print(f"Converted: {src} -> {dest}")
        except FileNotFoundError as e:
            print(f"Skipping {src} (tool missing?): {e}")
        except subprocess.CalledProcessError as e:
            print(f"LibreOffice failed on {src}: {e.stderr.decode(errors='ignore') if e.stderr else e}")
        except Exception as e:
            print(f"Error converting {src}: {e}")




if __name__ == "__main__":
    if __name__ == '__main__':
        input_directory = '/home/mh1f25/scratch/upcast_pro/odrl_extraction/row_data_processing_v2/data_sharing_agreement_2'
        output_directory = '/home/mh1f25/scratch/upcast_pro/odrl_extraction/row_data_processing_v2/data_sharing_agreement_2_text'
        convert_files_to_text(input_directory, output_directory)
